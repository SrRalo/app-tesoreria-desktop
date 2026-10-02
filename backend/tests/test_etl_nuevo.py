"""test_etl_nuevo.py — RF-ETL-01..14: normalizar, lectura, mapeo, preview, batch."""
from __future__ import annotations

import csv
import time

from openpyxl import Workbook

from etl.importar import importar, previsualizar
from etl.lectura import leer_archivo
from etl.mapeo import Catalogo
from etl.normalizar import es_fila_total, limpiar_texto, parse_fecha, parse_monto
from nucleo.basedatos import conectar

COLS = ["banco", "fecha_pago", "tipo", "tipo_pago", "entidad",
        "concepto_pago", "centro_costo", "valor_usd", "status", "observacion"]


def _xlsx(path, filas, titulos_extra=0):
    wb = Workbook()
    ws = wb.active
    ws.title = "Movimientos"
    for _ in range(titulos_extra):
        ws.append(["REPORTE DE TESORERIA", "", "", ""])
    ws.append(COLS)
    for f in filas:
        ws.append(f)
    wb.save(str(path))
    return path


BUENA = ["Pichincha", "2026-01-05", "ingreso", "transferencia", "CLI1",
         "prestamo", "", 8500, "realizado", "Cobro"]


# --- RF-ETL-04 fechas ---
def test_fecha_serial_excel_y_textos():
    iso, av = parse_fecha(46244)
    assert iso == "2026-08-10" and av is None
    assert parse_fecha("05/02/2026")[0] == "2026-02-05"
    assert parse_fecha("2026-08-14 00:00:00")[0] == "2026-08-14"
    iso, av = parse_fecha("no-fecha")
    assert iso is None and av  # va a revisión, no lanza


# --- RF-ETL-05 montos ---
def test_monto_simbolos_y_coma_decimal():
    assert parse_monto("$1,200")[0] == 1200.0
    assert parse_monto("USD 1.234,56")[0] == 1234.56
    assert parse_monto("(1,200)")[0] == -1200.0
    assert parse_monto(None, debito="200", credito="")[0] == -200.0
    assert parse_monto("abc")[0] == 0.0  # aviso, no lanza


# --- RF-ETL-06 texto + checklist totales ---
def test_texto_limpio_y_totales():
    assert limpiar_texto("  hola  mundo ", upper=True) == "HOLA MUNDO"
    assert es_fila_total(["TOTAL", "", 5000])
    assert not es_fila_total(["PACIFICCAM", "", 5000])


# --- RF-ETL-01/02 lectura ---
def test_header_dinamico_saltea_titulos(tmp_path):
    p = _xlsx(tmp_path / "t.xlsx", [BUENA], titulos_extra=2)
    filas, meta = leer_archivo(p)
    assert len(filas) == 1 and "header fila 3" in meta["header_info"]


def test_csv_con_punto_y_coma(tmp_path):
    p = tmp_path / "d.csv"
    with open(p, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(COLS)
        w.writerow(BUENA)
    filas, meta = leer_archivo(p)
    assert len(filas) == 1 and meta["formato"] == ".csv"


def test_xls_real_biff(tmp_path):
    xlwt = __import__("pytest").importorskip("xlwt")
    p = tmp_path / "d.xls"
    bk = xlwt.Workbook()
    sh = bk.add_sheet("Movimientos")
    for j, c in enumerate(COLS):
        sh.write(0, j, c)
    for j, v in enumerate(BUENA):
        sh.write(1, j, v)
    bk.save(str(p))
    filas, meta = leer_archivo(p)
    assert len(filas) == 1 and meta["formato"] == ".xls"


def test_debito_credito_separados(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Movimientos"
    ws.append(["banco", "fecha", "debito", "credito", "entidad"])
    ws.append(["Pichincha", "2026-01-05", 200, "", "Prov X"])
    p = tmp_path / "dc.xlsx"
    wb.save(str(p))
    filas, _ = leer_archivo(p)
    assert filas[0]["valor_usd"] == -200.0


# --- RF-ETL-07/08/09 mapeo ---
def test_fuzzy_y_fallback(con):
    with con:
        con.execute("INSERT INTO entidades (tipo, nombre) VALUES ('cliente','PACIFICCAM')")
    cat = Catalogo.cargar(con)
    cat.asegurar_fallbacks(con)
    from etl.mapeo import resolver_cuenta, resolver_entidad
    eid, av = resolver_entidad(cat, "pacificcam ", "cliente")  # exacto case-insens
    assert eid is not None and av is None
    eid2, av2 = resolver_entidad(cat, "PACIFICAM", "cliente")  # typo -> fuzzy
    assert eid2 == eid
    eid3, av3 = resolver_entidad(cat, "INEXISTENTE XYZ", "cliente")
    assert eid3 is None and "Por Definir" in av3
    cta, _ = resolver_cuenta(cat, "Pichincha CTA CTE 2100319432")  # alias+numero
    real = con.execute("SELECT id FROM cuentas WHERE banco='Pichincha'").fetchone()[0]
    assert cta == real


# --- RF-ETL-11 preview + RF-ETL-10/13/14 import ---
def test_preview_no_guarda_y_clasifica(tmp_path):
    p = _xlsx(tmp_path / "p.xlsx", [BUENA, ["Banco X", "mala-fecha", "ingreso",
             "transferencia", "Desconocido", "raro", "", "$500", "pendiente", ""]])
    db = tmp_path / "p.db"
    res = previsualizar(p, db)
    assert res["total"] == 2
    assert len(res["validas"]) + len(res["advertencias"]) == 2
    assert res["erroneas"] == []  # nada se descarta por fecha/mapeo
    con = conectar(db)
    try:
        assert con.execute("SELECT COUNT(*) c FROM movimientos").fetchone()["c"] == 0
    finally:
        con.close()


def test_import_batch_rapido_y_antipérdida(tmp_path):
    filas = [["Pichincha", "2026-01-05", "ingreso", "transferencia", f"CLI{i}",
              "prestamo", "", 100 + i, "realizado", ""] for i in range(300)]
    filas.append(["", "no-fecha", "ingreso", "", "", "", "", "USD 50", "", ""])  # todo malo
    p = _xlsx(tmp_path / "b.xlsx", filas)
    db = tmp_path / "b.db"
    t0 = time.time()
    res = importar(p, db)
    dt = time.time() - t0
    assert res["filas_ok"] == 301, res["errores"][:3]  # la fila mala también entra
    assert dt < 5.0, f"lote 300 filas tardó {dt:.1f}s"
    res2 = importar(p, db)  # reimport idempotente por dedup
    assert res2["filas_ok"] == 0 and res2["duplicadas"] == 301
