"""Extractos bancarios sin conciliación: solo carga a extracto_lineas + cortes.

Los estados de cuenta NO generan movimientos ni entran al Flujo; solo se
consultan en Entidades > Bancos. Parsers con fixtures sintéticas (sin
depender de los archivos reales); más humo con archivos reales si existen.
"""
from __future__ import annotations

from pathlib import Path

import pytest

from etl.extractos import internacional, pichincha, produbanco
from etl.extractos.base import es_comision, hash_linea, validar_cadena
from etl.extractos.importar_extracto import importar_extracto
from servicios import bancos as srv_bancos

PICH_HTML = """<html><head></head><body><table border="1">
<tr><td><b>Saldo Anterior:</b></td><td colspan="8">1000.00</td></tr>
<tr><td><b>Depositos/Creditos:</b></td><td colspan="8">500.00</td></tr>
<tr><td><b>Cheques/Débitos:</b></td><td colspan="8">300.00</td></tr>
<tr><td><b>Saldo Actual:</b></td><td colspan="8">1200.00</td></tr>
<tr><td><b>Cuenta:</b></td><td colspan="5"> CUENTA 2100319432 </td></tr>
<tr><td><b>Fecha este Corte:</b></td><td colspan="2">31-08-2026</td></tr>
<tr><td colspan="9">RESUMEN DE CHEQUES</td></tr>
<tr><td>999</td><td>03-ago.</td><td>300.00</td></tr>
<tr><td colspan="9">DETALLE DE MOVIMIENTOS</td></tr>
<tr><td>FECHA</td><td>OFIC.</td><td>N.DOC.</td><td>DESCRIPCION</td>
<td>DEBITO</td><td>CREDITO</td><td>SALDO</td></tr>
<tr><td>03-ago.</td><td>8386</td><td>DOC1</td><td>TRANSFERENCIA RECIBIDA</td>
<td>0.00</td><td>500.00</td><td>1500.00</td></tr>
<tr><td>04-ago.</td><td>0012</td><td>DOC2</td><td>PAGO PROVEEDOR</td>
<td>300.00</td><td>0.00</td><td>1200.00</td></tr>
</table></body></html>"""


def _mov(con, fecha, tipo, banco, valor, status="pendiente",
         concepto="pago_proveedores", entidad=None):
    con.execute(
        "INSERT INTO movimientos (fecha_pago, tipo, tipo_pago, concepto_id,"
        " entidad_id, valor_usd, status, observacion)"
        " VALUES (?,?, 'transferencia',"
        " (SELECT id FROM conceptos WHERE nombre=?),"
        " (SELECT id FROM entidades WHERE nombre=?), ?, ?, '')",
        (fecha, tipo, concepto, entidad, valor, status))
    return con.execute("SELECT last_insert_rowid()").fetchone()[0]


def _linea(con, banco, fecha, debito=0.0, credito=0.0, desc="X", ref="R"):
    cta = con.execute("SELECT id FROM cuentas WHERE banco=?", (banco,)).fetchone()[0]
    con.execute(
        "INSERT INTO extracto_lineas (cuenta_id, fecha, descripcion, referencia,"
        " debito_usd, credito_usd, hash_unico) VALUES (?,?,?,?,?,?,?)",
        (cta, fecha, desc, ref, debito, credito,
         hash_linea(cta, fecha, debito, credito, ref)))
    return con.execute("SELECT last_insert_rowid()").fetchone()[0]


class TestBase:
    def test_es_comision(self):
        assert es_comision("COM-11495538-LARVAS-PPV")
        assert es_comision("COSTO IVA CASH")
        assert es_comision("TARIFA TRANSF SCI RECIBIDA")
        assert not es_comision("TRANSFERENCIA RECIBIDA")
        assert not es_comision("PAGO CHEQ. VENTANILLA")

    def test_validar_cadena_ok_y_rota(self):
        lns = [{"fecha": "2026-08-03", "descripcion": "a",
                "debito": 0, "credito": 500, "saldo_banco": 1500},
               {"fecha": "2026-08-04", "descripcion": "b",
                "debito": 300, "credito": 0, "saldo_banco": 1200}]
        assert validar_cadena(lns, 1000) == []
        lns[1]["saldo_banco"] = 9999
        assert len(validar_cadena(lns, 1000)) == 1


class TestPichincha:
    def test_parsea_detalle_y_no_resumen(self, tmp_path):
        p = tmp_path / "pich.xls"
        p.write_text(PICH_HTML, encoding="utf-8")
        d = pichincha.parsear(p)
        assert (d["banco"], d["numero"]) == ("Pichincha", "2100319432")
        assert (d["saldo_anterior"], d["saldo_actual"]) == (1000, 1200)
        assert len(d["lineas"]) == 2  # resumen de cheques no se importa
        assert d["lineas"][0]["credito"] == 500
        assert validar_cadena(d["lineas"], d["saldo_anterior"]) == []


class TestProdubanco:
    def test_parsea_signos(self, tmp_path):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws["F11"] = "$  35.44"
        ws.cell(9, 6, "02006198356")
        # Posiciones reales del reporte: D Fecha, E Referencia, H Transacción,
        # I Signo, K Valor, N Saldo Contable.
        ws.cell(13, 4, "Fecha")
        ws.cell(13, 5, "Referencia")
        ws.cell(13, 8, "Transacción")
        ws.cell(13, 9, "Signo")
        ws.cell(13, 11, "Valor")
        ws.cell(13, 14, "Saldo Contable")
        ws.cell(14, 4, "2026-08-14 13:48:00")
        ws.cell(14, 5, "NOTA DE CREDITO")
        ws.cell(14, 8, "TRANSFERENCIA CTAS")
        ws.cell(14, 9, "(+)")
        ws.cell(14, 11, "$ 800.00")
        ws.cell(14, 14, "$ 46769.73")
        ws.cell(15, 4, "2026-08-18 11:49:00")
        ws.cell(15, 5, "PAGO DE CHEQUE")
        ws.cell(15, 8, "PAGO CHEQ.")
        ws.cell(15, 9, "(-)")
        ws.cell(15, 11, "$ 4000.00")
        ws.cell(15, 14, "$ 42769.73")
        p = tmp_path / "pro.xlsx"
        wb.save(str(p))
        d = produbanco.parsear(p)
        assert d["banco"] == "Produbanco"
        assert len(d["lineas"]) == 2
        assert d["lineas"][0]["credito"] == 800
        assert d["lineas"][1]["debito"] == 4000
        assert validar_cadena(d["lineas"], d["saldo_anterior"]) == []


class TestInternacional:
    def test_parsea_biff(self, tmp_path, monkeypatch):
        import etl.extractos.internacional as inter

        celdas = [
            ["", "Cuenta:", "7100614609"] + [""] * 8,
            ["", "Fecha", "Cod.", "Descripción", "", "Referencia Adicional",
             "", "Débito", "Crédito", "Saldo", "Ciudad"],
            ["", "2026-08-05", "TW", "TRF.PD MENIER S A", "", "",
             "", 0.0, 12573.0, 65690.95, "QUITO"],
            ["", "2026-08-06", "CQ", "CHEQUE POR CAMARA 217", "", "",
             "", 9536.0, 0.0, 56154.95, "QUITO"],
        ]

        class Celda:
            def __init__(self, v):
                self.value = v

        class Hoja:
            nrows, ncols = 4, 11

            def cell(self, r, c):
                return Celda(celdas[r][c])

        class Libro:
            def sheet_by_index(self, i):
                return Hoja()

        monkeypatch.setattr("xlrd.open_workbook",
                            lambda *a, **k: Libro())
        d = inter.parsear(tmp_path / "inter.xls")
        assert d["banco"] == "Internacional"
        assert len(d["lineas"]) == 2
        assert d["saldo_anterior"] == pytest.approx(53117.95)
        assert validar_cadena(d["lineas"], d["saldo_anterior"]) == []


class TestSinConciliacion:
    def test_importar_no_genera_movimientos(self, tmp_path, db):
        p = tmp_path / "pich.xls"
        p.write_text(PICH_HTML, encoding="utf-8")
        antes = None
        from nucleo.basedatos import conectar
        con = conectar(db)
        try:
            antes = con.execute("SELECT COUNT(*) c FROM movimientos").fetchone()["c"]
        finally:
            con.close()
        r = importar_extracto(p, "pichincha", db, nombre_original="pich.xls")
        assert r["lineas_nuevas"] == 2 and r["apertura_fijada"] == 1000
        assert "conciliacion" not in r  # sin conciliación
        con = conectar(db)
        try:
            despues = con.execute("SELECT COUNT(*) c FROM movimientos").fetchone()["c"]
            assert despues == antes  # el extracto no entra al flujo
            assert con.execute("SELECT COUNT(*) c FROM conciliacion").fetchone()["c"] == 0
            assert con.execute("SELECT COUNT(*) c FROM extracto_lineas").fetchone()["c"] == 2
            # Visible solo en Bancos
            bancos = srv_bancos.listar_bancos(con)
            assert len(bancos) == 1 and bancos[0]["banco"] == "Pichincha"
            assert bancos[0]["saldo"] == 1200
        finally:
            con.close()

    def test_importar_extracto_idempotente(self, tmp_path, db):
        p = tmp_path / "pich.xls"
        p.write_text(PICH_HTML, encoding="utf-8")
        r1 = importar_extracto(p, "pichincha", db, nombre_original="pich.xls")
        assert r1["lineas_nuevas"] == 2 and r1["apertura_fijada"] == 1000
        r2 = importar_extracto(p, "pichincha", db, nombre_original="pich.xls")
        assert r2["lineas_nuevas"] == 0 and r2["duplicadas"] == 2
        assert r2["apertura_fijada"] is None  # la apertura no se pisa

    @pytest.mark.parametrize("ruta", [
        Path(r"C:\Users\Carlos\Desktop\practicas 2026-2\fuente de datos")
        / "01 31 agosto Pichiccha.xls"])
    def test_humo_archivo_real_pichincha(self, ruta):
        if not ruta.exists():
            pytest.skip("sin acceso a la ruta de datos")
        from etl.extractos.importar_extracto import parsear_extracto
        d = parsear_extracto(ruta, "pichincha")
        assert d["saldo_actual"] == pytest.approx(236014.39)
        assert not validar_cadena(d["lineas"], d["saldo_anterior"])
