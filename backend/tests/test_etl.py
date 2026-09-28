"""Fase 4 — etl/importar.py + etl/plantilla.py: ETL y plantilla paso a paso."""
from __future__ import annotations

import pytest
from openpyxl import Workbook, load_workbook

from etl.importar import COLUMNAS, importar, leer_filas, validar
from etl.plantilla import crear_plantilla
from nucleo.basedatos import conectar

BUENA = ["Pichincha", "2026-01-05", "ingreso", "transferencia", "CLI1",
         "prestamo", "", 8500, "realizado", "Cobro"]


def xlsx_con_filas(path, filas):
    wb = Workbook()
    ws = wb.active
    ws.title = "Movimientos"
    ws.append(COLUMNAS)
    for f in filas:
        ws.append(f)
    wb.save(str(path))
    return path


def test_plantilla_trae_formato_y_desplegables(tmp_path):
    out = crear_plantilla(tmp_path / "plantilla_flujo.xlsx")
    wb = load_workbook(str(out))
    assert "Movimientos" in wb.sheetnames and "Ayuda" in wb.sheetnames
    ws = wb["Movimientos"]
    assert [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))] == COLUMNAS
    assert len(list(ws.data_validations.dataValidation)) >= 4  # tipo, tipo_pago, concepto, status


def test_fila_invalida_se_reporta_sin_romper_carga(tmp_path):
    mala = ["Pichincha", "no-fecha", "otro", "transferencia", "", "nomina", "", -5,
            "raro", ""]
    xlsx = xlsx_con_filas(tmp_path / "mix.xlsx", [BUENA, mala])
    db = tmp_path / "mix.db"
    res = importar(xlsx, db)
    assert res["filas_ok"] == 1
    assert any("fila 3" in e for e in res["errores"])
    con = conectar(db)
    try:
        assert con.execute("SELECT COUNT(*) c FROM movimientos").fetchone()["c"] == 1
    finally:
        con.close()


def test_acepta_fecha_latina_y_valor_con_formato(tmp_path):
    fila = ["Caja", "05/02/2026", "egreso", "efectivo", "", "nomina", "", "$1,200",
            "pendiente", ""]
    xlsx = xlsx_con_filas(tmp_path / "lat.xlsx", [fila])
    db = tmp_path / "lat.db"
    assert importar(xlsx, db)["filas_ok"] == 1
    con = conectar(db)
    try:
        row = con.execute("SELECT fecha_pago, valor_usd FROM movimientos").fetchone()
        assert (row["fecha_pago"], row["valor_usd"]) == ("2026-02-05", 1200)
    finally:
        con.close()


def test_etl_crea_concepto_cuenta_y_entidad_nuevos(tmp_path):
    fila = ["Banco X", "2026-01-05", "ingreso", "cheque", "Nuevo CLI", "servicios",
            "", 100, "realizado", ""]
    xlsx = xlsx_con_filas(tmp_path / "nuevo.xlsx", [fila])
    db = tmp_path / "nuevo.db"
    assert importar(xlsx, db)["filas_ok"] == 1
    con = conectar(db)
    try:
        assert con.execute(
            "SELECT 1 FROM conceptos WHERE nombre='servicios'").fetchone()
        assert con.execute(
            "SELECT 1 FROM cuentas WHERE banco='Banco X'").fetchone()
        assert con.execute(
            "SELECT tipo FROM entidades WHERE nombre='Nuevo CLI'").fetchone()["tipo"] == "cliente"
    finally:
        con.close()


def test_sin_hoja_movimientos_aborta(tmp_path):
    wb = Workbook()
    wb.active.title = "Otra"
    p = tmp_path / "mala.xlsx"
    wb.save(str(p))
    with pytest.raises(SystemExit):
        leer_filas(p)


def test_validar_detecta_todo(tmp_path):
    d = dict(zip(COLUMNAS, ["", "x", "otro", "tarjeta", "", "", "", 0, "raro", ""]))
    d["_fila"] = 9
    errs = validar(d)
    assert len(errs) >= 5 and all("fila 9" in e for e in errs)
